import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    # 16:9 widescreen
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Theme Colors
    BG_DARK = RGBColor(10, 15, 29)       # #0A0F1D Deep Navy
    CARD_BG = RGBColor(19, 27, 46)       # #131B2E Card Navy
    CARD_BORDER = RGBColor(30, 43, 74)   # #1E2B4A
    CYAN = RGBColor(0, 229, 255)         # #00E5FF Accent Cyan
    BLUE = RGBColor(0, 163, 255)         # #00A3FF Electric Blue
    GREEN = RGBColor(0, 230, 118)        # #00E676 Emerald Green
    AMBER = RGBColor(255, 179, 0)        # #FFB300 Warning Amber
    RED = RGBColor(255, 61, 87)          # #FF3D57 Crimson Threat
    WHITE = RGBColor(255, 255, 255)
    GRAY = RGBColor(148, 163, 184)       # #94A3B8 Slate Gray
    LIGHT_GRAY = RGBColor(203, 213, 225) # #CBD5E1

    def set_bg(slide):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = BG_DARK

    def add_header(slide, title_text, category_text="SHIELD AI // THREAT INTELLIGENCE"):
        # Category / Tracker Pill
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.4))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.size = Pt(10)
        p_cat.font.bold = True
        p_cat.font.color.rgb = CYAN

        # Main Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.7), Inches(0.8))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(24)
        p_title.font.bold = True
        p_title.font.color.rgb = WHITE

    def add_card(slide, left, top, width, height, title="", border_color=CARD_BORDER):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = border_color
        card.line.width = Pt(1.5)
        return card

    # ==========================================
    # SLIDE 1: Title Slide
    # ==========================================
    slide_layout = prs.slide_layouts[6] # Blank
    s1 = prs.slides.add_slide(slide_layout)
    set_bg(s1)

    # Decorative Cyan Line at top
    dec_line = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.2), Inches(1.2), Inches(0.06))
    dec_line.fill.solid()
    dec_line.fill.fore_color.rgb = CYAN
    dec_line.line.fill.background()

    # Title Box
    t_box = s1.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11.7), Inches(2.2))
    tf1 = t_box.text_frame
    tf1.word_wrap = True
    p1 = tf1.paragraphs[0]
    p1.text = "SHIELD AI"
    p1.font.size = Pt(48)
    p1.font.bold = True
    p1.font.color.rgb = WHITE

    p1_sub = tf1.add_paragraph()
    p1_sub.text = "Next-Generation In-Browser Phishing, Quishing & Homoglyph Countermeasure System"
    p1_sub.font.size = Pt(20)
    p1_sub.font.color.rgb = CYAN
    p1_sub.space_before = Pt(12)

    # Feature badges card
    add_card(s1, Inches(0.8), Inches(4.0), Inches(11.7), Inches(2.5))
    badge_box = s1.shapes.add_textbox(Inches(1.1), Inches(4.2), Inches(11.1), Inches(2.1))
    bf = badge_box.text_frame
    bf.word_wrap = True

    bp0 = bf.paragraphs[0]
    bp0.text = "CORE CAPABILITIES AT A GLANCE"
    bp0.font.size = Pt(11)
    bp0.font.bold = True
    bp0.font.color.rgb = BLUE

    capabilities = [
        "• On-Demand 'Squeeze' Sidebar: Sider AI-inspired physical viewport compression calc(100% - 420px)",
        "• Interactive Snipping Tool: Canvas-cropped bounding box analysis powered by Gemini 3.5/3.8 Flash Vision",
        "• Quishing Matrix Defense: Direct 2D QR decoding (jsQR + Jimp) with recursive 10-hop redirect tracing",
        "• Brand Typosquatting Engine: Detection of homoglyphs like 'rnicrosoft.com' with enforced 98/100 risk floors",
        "• Active Countermeasures: Autonomous Scam-Baiter Bot, Ask Shield Copilot & Forensic SOC Incident PDF Export"
    ]
    for cap in capabilities:
        bp = bf.add_paragraph()
        bp.text = cap
        bp.font.size = Pt(13)
        bp.font.color.rgb = LIGHT_GRAY
        bp.space_before = Pt(5)

    # Footer
    f_box = s1.shapes.add_textbox(Inches(0.8), Inches(6.7), Inches(11.7), Inches(0.4))
    ff = f_box.text_frame
    fp = ff.paragraphs[0]
    fp.text = "TCET HACKATHON 2026  |  ENTERPRISE & CONSUMER CYBERSECURITY TRACK"
    fp.font.size = Pt(10)
    fp.font.color.rgb = GRAY

    # ==========================================
    # SLIDE 2: The Critical Problem
    # ==========================================
    s2 = prs.slides.add_slide(slide_layout)
    set_bg(s2)
    add_header(s2, "The Evolving Threat Landscape: Why Traditional Antivirus Fails", "01 // THE THREAT CONTEXT")

    # 3 Stat Cards
    stat_data = [
        ("3.4 Billion", "Daily Phishing Messages", "Attackers now use LLMs to clone brand styles with flawless grammar and persuasive context.", RED, Inches(0.8)),
        ("78% Spurt", "Quishing (QR Phishing)", "Attackers embed malicious URLs inside QR codes in emails and flyers, completely bypassing text firewalls.", AMBER, Inches(4.8)),
        ("0-Day Domains", "Homoglyph Lookalikes", "Deceptive domains like 'rnicrosoft.com' (rn -> m) exploit visual perception before blacklists can detect them.", CYAN, Inches(8.8))
    ]

    for val, title, desc, col, x_pos in stat_data:
        add_card(s2, x_pos, Inches(1.8), Inches(3.7), Inches(4.8), border_color=col)
        box = s2.shapes.add_textbox(x_pos + Inches(0.2), Inches(2.0), Inches(3.3), Inches(4.4))
        tf = box.text_frame
        tf.word_wrap = True

        p_val = tf.paragraphs[0]
        p_val.text = val
        p_val.font.size = Pt(32)
        p_val.font.bold = True
        p_val.font.color.rgb = col

        p_t = tf.add_paragraph()
        p_t.text = title
        p_t.font.size = Pt(16)
        p_t.font.bold = True
        p_t.font.color.rgb = WHITE
        p_t.space_before = Pt(10)

        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(13)
        p_d.font.color.rgb = LIGHT_GRAY
        p_d.space_before = Pt(12)

        p_sub = tf.add_paragraph()
        p_sub.text = "FLAW OF LEGACY TOOLS:"
        p_sub.font.size = Pt(11)
        p_sub.font.bold = True
        p_sub.font.color.rgb = GRAY
        p_sub.space_before = Pt(20)

        p_sub_desc = tf.add_paragraph()
        p_sub_desc.text = "Static database lookups have an 8-48 hour delay. Attackers switch infrastructure in minutes."
        p_sub_desc.font.size = Pt(12)
        p_sub_desc.font.color.rgb = GRAY
        p_sub_desc.space_before = Pt(4)

    # ==========================================
    # SLIDE 3: The Solution — Shield AI
    # ==========================================
    s3 = prs.slides.add_slide(slide_layout)
    set_bg(s3)
    add_header(s3, "The Solution: A Real-Time, Multimodal In-Browser Defense Agent", "02 // PRODUCT VISION")

    sol_cards = [
        ("⚡ Non-Intrusive Zero-Friction UI", "Squeeze Sidebar compresses host website viewport margin without popups or obscuring page controls.", CYAN, Inches(0.8), Inches(1.8)),
        ("👁️ Multimodal Visual OCR", "Interactive Snipping Tool feeds cropped sub-images directly to Gemini Vision to detect fake branding & urgent traps.", BLUE, Inches(6.8), Inches(1.8)),
        ("🔍 Recursive Redirect Traversal", "Traces shortlinks (bit.ly, tinyurl) up to 10 HTTP HEAD hops deep to find the unmasked destination.", GREEN, Inches(0.8), Inches(4.3)),
        ("🤖 Active Countermeasures & Recovery", "Autonomous Scam-Baiter wastes attacker resources; Recovery Wizard gives step-by-step account hardening.", AMBER, Inches(6.8), Inches(4.3))
    ]

    for title, desc, col, x, y in sol_cards:
        add_card(s3, x, y, Inches(5.7), Inches(2.2), border_color=col)
        box = s3.shapes.add_textbox(x + Inches(0.3), y + Inches(0.2), Inches(5.1), Inches(1.8))
        tf = box.text_frame
        tf.word_wrap = True

        p_t = tf.paragraphs[0]
        p_t.text = title
        p_t.font.size = Pt(17)
        p_t.font.bold = True
        p_t.font.color.rgb = col

        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(13)
        p_d.font.color.rgb = LIGHT_GRAY
        p_d.space_before = Pt(8)

    # ==========================================
    # SLIDE 4: Squeeze Sidebar UI
    # ==========================================
    s4 = prs.slides.add_slide(slide_layout)
    set_bg(s4)
    add_header(s4, "Engineering Innovation: On-Demand 'Squeeze' Sidebar", "03 // FRONTEND ARCHITECTURE")

    # Left: Explanation
    add_card(s4, Inches(0.8), Inches(1.8), Inches(5.7), Inches(4.8))
    box_l = s4.shapes.add_textbox(Inches(1.1), Inches(2.0), Inches(5.1), Inches(4.4))
    tfl = box_l.text_frame
    tfl.word_wrap = True

    p = tfl.paragraphs[0]
    p.text = "Why Most Security Extensions Get Uninstalled"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = WHITE

    points = [
        "❌ Invasive Popups: Float over input forms, closing whenever the user clicks outside.",
        "❌ Screen Distortion: Standard modal alerts lock the entire tab and frustrate users.",
        "✅ The Squeeze Model: Modeled after Sider AI & REC Note. Dynamically compresses the host page.",
        "• CSS Viewport Manipulation: Toggles width: calc(100% - 420px) and margin-right: 420px.",
        "• 0.28s Smooth Cubic-Bezier Transition: Zero jarring layout jumps.",
        "• Complete Iframe Sandboxing: The sidebar is isolated from host DOM scripts and malicious tampering."
    ]
    for pt in points:
        bp = tfl.add_paragraph()
        bp.text = pt
        bp.font.size = Pt(13)
        bp.font.color.rgb = CYAN if "✅" in pt else LIGHT_GRAY
        bp.space_before = Pt(8)

    # Right: Visual Comparison Card
    add_card(s4, Inches(6.8), Inches(1.8), Inches(5.7), Inches(4.8), border_color=BLUE)
    box_r = s4.shapes.add_textbox(Inches(7.1), Inches(2.0), Inches(5.1), Inches(4.4))
    tfr = box_r.text_frame
    tfr.word_wrap = True

    p = tfr.paragraphs[0]
    p.text = "Sidebar Component Layout"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = BLUE

    diagram_text = (
        "+--------------------------------+-------------------+\n"
        "| HOST WEBPAGE                   | SHIELD SIDEBAR    |\n"
        "| (calc(100% - 420px))           | (420px fixed)     |\n"
        "|                                |                   |\n"
        "|  Forms, buttons, and links     | [ Auto-Scan: ON ] |\n"
        "|  remain 100% clickable while   | [ Scan Area (Snip)]\n"
        "|  Shield sits beside them       |                   |\n"
        "|                                | VERDICT: 98/100   |\n"
        "|  [ Suspicious Login Box ]      | [ BLOCKED: RED ]  |\n"
        "|                                |                   |\n"
        "|                                | Ask Copilot / PDF |\n"
        "+--------------------------------+-------------------+"
    )
    bp = tfr.add_paragraph()
    bp.text = diagram_text
    bp.font.size = Pt(11)
    bp.font.name = "Courier New"
    bp.font.color.rgb = CYAN
    bp.space_before = Pt(10)

    # ==========================================
    # SLIDE 5: Interactive Snipping Tool & Vision
    # ==========================================
    s5 = prs.slides.add_slide(slide_layout)
    set_bg(s5)
    add_header(s5, "Interactive Snipping Tool: Multimodal Vision at the Click of a Button", "04 // COMPUTER VISION")

    snip_steps = [
        ("1. User Trigger", "Clicks 'Scan Area' in the sidebar. The host page dims with crosshair cursor listeners.", Inches(0.8)),
        ("2. Canvas Bounding Box", "User drags over suspicious form. Extension uses captureVisibleTab + offscreen canvas to slice the exact rectangle.", Inches(3.8)),
        ("3. Context Extraction", "Simultaneously harvests surrounding DOM text nodes, button captions, and form input names.", Inches(6.8)),
        ("4. Gemini Vision OCR", "Multimodal LLM verifies logo authenticity, deceptive countdowns, and unverified credential harvesting.", Inches(9.8))
    ]

    for title, desc, x in snip_steps:
        add_card(s5, x, Inches(1.8), Inches(2.7), Inches(4.8), border_color=CYAN)
        box = s5.shapes.add_textbox(x + Inches(0.2), Inches(2.0), Inches(2.3), Inches(4.4))
        tf = box.text_frame
        tf.word_wrap = True

        p_t = tf.paragraphs[0]
        p_t.text = title
        p_t.font.size = Pt(16)
        p_t.font.bold = True
        p_t.font.color.rgb = CYAN

        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(13)
        p_d.font.color.rgb = LIGHT_GRAY
        p_d.space_before = Pt(12)

    # ==========================================
    # SLIDE 6: Quishing & Homoglyph Defense
    # ==========================================
    s6 = prs.slides.add_slide(slide_layout)
    set_bg(s6)
    add_header(s6, "Stopping Quishing (QR Attacks) & Brand Homoglyphs", "05 // THREAT ENGINE")

    # Left: Quishing
    add_card(s6, Inches(0.8), Inches(1.8), Inches(5.7), Inches(4.8), border_color=AMBER)
    box_q = s6.shapes.add_textbox(Inches(1.1), Inches(2.0), Inches(5.1), Inches(4.4))
    tfq = box_q.text_frame
    tfq.word_wrap = True

    pq = tfq.paragraphs[0]
    pq.text = "Quishing Matrix Engine (QR Phishing)"
    pq.font.size = Pt(18)
    pq.font.bold = True
    pq.font.color.rgb = AMBER

    q_points = [
        "• Problem: QR codes conceal malicious targets from standard text filters.",
        "• Direct Bitmap Decoding: Jimp extracts RGBA pixel array; jsQR analyzes finder patterns and unmasks data bits.",
        "• 10-Hop Redirect Unraveling: Immediately follows shorteners to the final payload server.",
        "• High Precision Distinction:",
        "    - Legitimate menus/support: 0 / 100 (Safe)",
        "    - KYC / credential harvesting QR: 95-98 / 100 (Critical Threat)"
    ]
    for pt in q_points:
        bp = tfq.add_paragraph()
        bp.text = pt
        bp.font.size = Pt(13)
        bp.font.color.rgb = LIGHT_GRAY
        bp.space_before = Pt(8)

    # Right: Homoglyphs
    add_card(s6, Inches(6.8), Inches(1.8), Inches(5.7), Inches(4.8), border_color=RED)
    box_h = s6.shapes.add_textbox(Inches(7.1), Inches(2.0), Inches(5.1), Inches(4.4))
    tfh = box_h.text_frame
    tfh.word_wrap = True

    ph = tfh.paragraphs[0]
    ph.text = "Brand Homoglyph Engine (Typosquatting)"
    ph.font.size = Pt(18)
    ph.font.bold = True
    ph.font.color.rgb = RED

    h_points = [
        "• The rn -> m Attack: 'www.rnicrosoft.com' appears identical to Microsoft on mobile & desktop screens.",
        "• Regex Normalization Pipeline:",
        "    - rn -> m (rnicrosoft -> microsoft)",
        "    - 0 -> o (micros0ft -> microsoft)",
        "    - 1 -> l (goog1e -> google, paypa1 -> paypal)",
        "    - Punycode detection (xn--)",
        "• Enforced Evidence Floor: Detected typosquats are locked at 98 / 100 Risk, bypassing all attacker evasion."
    ]
    for pt in h_points:
        bp = tfh.add_paragraph()
        bp.text = pt
        bp.font.size = Pt(13)
        bp.font.color.rgb = LIGHT_GRAY
        bp.space_before = Pt(8)

    # ==========================================
    # SLIDE 7: System Architecture & Workflow (With Image!)
    # ==========================================
    s7 = prs.slides.add_slide(slide_layout)
    set_bg(s7)
    add_header(s7, "End-to-End System Workflow Architecture", "06 // SYSTEM DIAGRAM")

    img_path = "c:\\Users\\AASHAY KULAKRNI\\Downloads\\TCET_Hackathon\\TCET_Hackathon\\TCET-hackathon\\user_workflow_diagram.jpg"
    if os.path.exists(img_path):
        s7.shapes.add_picture(img_path, Inches(0.8), Inches(1.6), Inches(11.73), Inches(5.3))
    else:
        # Fallback card if image missing
        add_card(s7, Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0))
        box_f = s7.shapes.add_textbox(Inches(1.1), Inches(2.0), Inches(11.1), Inches(4.6))
        box_f.text_frame.paragraphs[0].text = "Workflow Diagram (Image: user_workflow_diagram.jpg)"

    # ==========================================
    # SLIDE 8: Active Countermeasures & Post-Incident
    # ==========================================
    s8 = prs.slides.add_slide(slide_layout)
    set_bg(s8)
    add_header(s8, "Beyond Detection: Active Countermeasures & Incident Response", "07 // USER DEFENSE")

    actions = [
        ("Ask Shield Copilot", "Context-aware AI assistant answers security questions in plain, reassuring language with multi-model failover (Gemini -> Groq -> Local Heuristics).", CYAN, Inches(0.8), Inches(1.8)),
        ("Autonomous Scam-Baiter Bot", "Deploys 3 personas (Technical Glitch, Bureaucracy, Confused Customer) generating realistic friction to waste scammers' bandwidth safely.", BLUE, Inches(6.8), Inches(1.8)),
        ("Account Recovery Wizard", "Interactive 4-step checklist (Session Kill -> Official Password Reset -> Hardware 2FA -> Device Audit) with progress synced to database.", GREEN, Inches(0.8), Inches(4.3)),
        ("Forensic SOC PDF Report", "Exports print-ready incident documentation via jsPDF containing timestamps, 10-hop redirect chains, and AV consensus logs for IT teams.", RED, Inches(6.8), Inches(4.3))
    ]

    for title, desc, col, x, y in actions:
        add_card(s8, x, y, Inches(5.7), Inches(2.2), border_color=col)
        box = s8.shapes.add_textbox(x + Inches(0.3), y + Inches(0.2), Inches(5.1), Inches(1.8))
        tf = box.text_frame
        tf.word_wrap = True

        p_t = tf.paragraphs[0]
        p_t.text = title
        p_t.font.size = Pt(17)
        p_t.font.bold = True
        p_t.font.color.rgb = col

        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(13)
        p_d.font.color.rgb = LIGHT_GRAY
        p_d.space_before = Pt(8)

    # ==========================================
    # SLIDE 9: Technical Benchmarks & Verification
    # ==========================================
    s9 = prs.slides.add_slide(slide_layout)
    set_bg(s9)
    add_header(s9, "Empirical Validation: Rigorous 100% Pass Benchmarks", "08 // VERIFICATION")

    benchmarks = [
        ("15 / 15", "Visual Threat Benchmark", "scripts/verify-15-images.js\n• 5 Safe: 0/100 (Google/Gemini in Gmail)\n• 5 Suspicious: 45-60/100 (Urgent surveys)\n• 5 Dangerous: 95-98/100 (rnicrosoft.com, KYC)", GREEN, Inches(0.8)),
        ("8 / 8", "QR Matrix Decoder Suite", "scripts/verify-qr-suite.js\n• Tested authentic 2D PNG bitmaps\n• Distinguishes legitimate support links from quishing payloads with 100% precision", CYAN, Inches(4.8)),
        ("19 / 19", "Automated API Unit Tests", "test/api.test.js\n• Express endpoints\n• Redirect recursion limiters\n• Typosquatting regex replacements\n• Supabase RLS privacy isolation", BLUE, Inches(8.8))
    ]

    for score, title, desc, col, x in benchmarks:
        add_card(s9, x, Inches(1.8), Inches(3.7), Inches(4.8), border_color=col)
        box = s9.shapes.add_textbox(x + Inches(0.2), Inches(2.0), Inches(3.3), Inches(4.4))
        tf = box.text_frame
        tf.word_wrap = True

        p_s = tf.paragraphs[0]
        p_s.text = score
        p_s.font.size = Pt(38)
        p_s.font.bold = True
        p_s.font.color.rgb = col

        p_t = tf.add_paragraph()
        p_t.text = title
        p_t.font.size = Pt(16)
        p_t.font.bold = True
        p_t.font.color.rgb = WHITE
        p_t.space_before = Pt(8)

        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(13)
        p_d.font.color.rgb = LIGHT_GRAY
        p_d.space_before = Pt(12)

    # ==========================================
    # SLIDE 10: Team Roles & Ownership
    # ==========================================
    s10 = prs.slides.add_slide(slide_layout)
    set_bg(s10)
    add_header(s10, "Engineering Team: 4 Specialized Pillars", "09 // TEAM ROLES")

    roles = [
        ("Member 1", "Team Lead & System Architect", "Architecture & Layout Engine", [
            "• 2026 Phishing Crisis story",
            "• Squeeze Viewport Engine",
            "• Iframe Sandbox Isolation",
            "• Live demo orchestration"
        ], GREEN, Inches(0.7)),
        ("Member 2", "Extension & Vision Lead", "Browser Runtime & OCR", [
            "• Manifest V3 Service Worker",
            "• Interactive Snipping Tool",
            "• Offscreen Canvas sub-crop",
            "• Gemini 3.5/3.8 Flash Vision"
        ], CYAN, Inches(3.7)),
        ("Member 3", "Threat Intel & Networks", "Telemetry & Homoglyphs", [
            "• 10-hop redirect tracer",
            "• VirusTotal 70+ AV engines",
            "• Homoglyphs (rn -> m)",
            "• Evidence floor (98/100)"
        ], BLUE, Inches(6.7)),
        ("Member 4", "Defense & Incident Response", "Quishing, RLS & Forensics", [
            "• Quishing QR Matrix Decoder",
            "• Supabase PostgreSQL RLS",
            "• Autonomous Scam-Baiter Bot",
            "• Forensic SOC PDF Export"
        ], AMBER, Inches(9.7))
    ]

    for mem, title, subtitle, bullets, col, x in roles:
        add_card(s10, x, Inches(1.8), Inches(2.8), Inches(4.8), border_color=col)
        box = s10.shapes.add_textbox(x + Inches(0.15), Inches(2.0), Inches(2.5), Inches(4.4))
        tf = box.text_frame
        tf.word_wrap = True

        p_m = tf.paragraphs[0]
        p_m.text = mem.upper()
        p_m.font.size = Pt(13)
        p_m.font.bold = True
        p_m.font.color.rgb = col

        p_t = tf.add_paragraph()
        p_t.text = title
        p_t.font.size = Pt(16)
        p_t.font.bold = True
        p_t.font.color.rgb = WHITE
        p_t.space_before = Pt(4)

        p_s = tf.add_paragraph()
        p_s.text = subtitle
        p_s.font.size = Pt(11)
        p_s.font.color.rgb = GRAY
        p_s.space_before = Pt(2)

        for b in bullets:
            pb = tf.add_paragraph()
            pb.text = b
            pb.font.size = Pt(12)
            pb.font.color.rgb = LIGHT_GRAY
            pb.space_before = Pt(8)

    # ==========================================
    # SLIDE 11: Summary & Conclusion
    # ==========================================
    s11 = prs.slides.add_slide(slide_layout)
    set_bg(s11)
    add_header(s11, "The Future of In-Browser Cybersecurity", "10 // CONCLUSION & ROADMAP")

    # Left: Value summary
    add_card(s11, Inches(0.8), Inches(1.8), Inches(5.7), Inches(4.8), border_color=CYAN)
    box_vl = s11.shapes.add_textbox(Inches(1.1), Inches(2.0), Inches(5.1), Inches(4.4))
    tfvl = box_vl.text_frame
    tfvl.word_wrap = True

    p = tfvl.paragraphs[0]
    p.text = "Why Shield AI Wins in the Real World"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = CYAN

    wins = [
        "🛡️ Zero Delay: Real-time visual OCR & redirect unraveling catches zero-day threats before blacklists index them.",
        "✨ Frictionless UX: Physical 'Squeeze' sidebar doesn't block work or get uninstalled.",
        "🎯 High Precision: 100% accuracy on brand homoglyphs (rn -> m) and QR code quishing.",
        "🔒 Total Privacy: Supabase Row-Level Security guarantees users only see their own scans.",
        "⚡ Turnkey Deployment: Ready on Chrome Web Store & Vercel."
    ]
    for w in wins:
        bp = tfvl.add_paragraph()
        bp.text = w
        bp.font.size = Pt(13)
        bp.font.color.rgb = LIGHT_GRAY
        bp.space_before = Pt(10)

    # Right: Thank You & Live Demo
    add_card(s11, Inches(6.8), Inches(1.8), Inches(5.7), Inches(4.8), border_color=BLUE)
    box_vr = s11.shapes.add_textbox(Inches(7.1), Inches(2.2), Inches(5.1), Inches(4.0))
    tfvr = box_vr.text_frame
    tfvr.word_wrap = True

    p = tfvr.paragraphs[0]
    p.text = "READY FOR LIVE DEMONSTRATION"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = BLUE

    p_big = tfvr.add_paragraph()
    p_big.text = "Thank You, Judges!"
    p_big.font.size = Pt(36)
    p_big.font.bold = True
    p_big.font.color.rgb = WHITE
    p_big.space_before = Pt(10)

    p_repo = tfvr.add_paragraph()
    p_repo.text = "GitHub: github.com/AAKU2107/TCET-hackathon\nLive Web Console & Chrome Manifest V3 Extension Ready"
    p_repo.font.size = Pt(14)
    p_repo.font.color.rgb = CYAN
    p_repo.space_before = Pt(16)

    p_qa = tfvr.add_paragraph()
    p_qa.text = "We are excited to take your questions."
    p_qa.font.size = Pt(14)
    p_qa.font.color.rgb = LIGHT_GRAY
    p_qa.space_before = Pt(12)

    output_path = "c:\\Users\\AASHAY KULAKRNI\\Downloads\\TCET_Hackathon\\TCET_Hackathon\\TCET-hackathon\\Shield_AI_Pitch_Deck.pptx"
    try:
        prs.save(output_path)
        print(f"Presentation successfully saved to: {output_path}")
    except PermissionError:
        alt_path = "c:\\Users\\AASHAY KULAKRNI\\Downloads\\TCET_Hackathon\\TCET_Hackathon\\TCET-hackathon\\Shield_AI_Pitch_Deck_v2.pptx"
        prs.save(alt_path)
        print(f"Primary file was open in PowerPoint. Saved updated 4-member deck to: {alt_path}")

if __name__ == "__main__":
    create_deck()
